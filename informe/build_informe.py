"""Genera informe/informe.docx (APA 7, trabajo de estudiante) a partir de outputs/ y verifica su formato.

Ejecutar después del notebook: python informe/build_informe.py
Reglas: Arial 11, interlineado doble, márgenes de 2,54 cm, texto en negro, negrita solo en títulos y rótulos,
sin guiones largos ni medios.
"""
import datetime as dt
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
DOCX = Path(__file__).with_name("informe.docx")

FONT, SIZE, BLACK = "Arial", Pt(11), RGBColor(0, 0, 0)
FORBIDDEN = ("—", "–")  # guion largo y guion medio

TITLE = ("Detección temprana de anomalías en aerogeneradores con datos SCADA mediante modelos de "
         "comportamiento normal")
AUTHOR = "Nicolás Meneses Marín"
UNIVERSITY = "Universidad Pontificia Bolivariana"
COURSE = "Integración de Energías Renovables No Convencionales a Sistemas de Potencia"
TEACHERS = ["Idi A. Isaac Millan", "Jorge W. Gonzalez Sanchez", "Carlos E. Restrepo Mejia"]
MONTHS = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
          "noviembre", "diciembre"]


# ---------------------------------------------------------------- utilidades de formato
def fmt(x, d=1):
    """Número con coma decimal."""
    return f"{x:.{d}f}".replace(".", ",")


def set_font(target, bold=None, italic=None):
    """Fija Arial 11 negro en un run o estilo, incluidas las variantes eastAsia/cs y sin fuentes de tema."""
    f = target.font
    f.name, f.size, f.color.rgb = FONT, SIZE, BLACK
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    rpr = target._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), FONT)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{attr}"), None)


def para_format(p, indent=True, align=None, spacing=WD_LINE_SPACING.DOUBLE, hanging=False, keep=False):
    pf = p.paragraph_format
    pf.line_spacing_rule = spacing
    pf.space_before = pf.space_after = Pt(0)
    pf.first_line_indent = Inches(0.5) if indent else Inches(0)
    if hanging:
        pf.left_indent, pf.first_line_indent = Inches(0.5), Inches(-0.5)
    if align is not None:
        p.alignment = align
    pf.keep_with_next = keep
    pf.widow_control = True
    return p


def add_runs(p, text):
    """Texto con *cursiva* opcional (solo donde APA la exige)."""
    for i, chunk in enumerate(re.split(r"\*", text)):
        if chunk:
            set_font(p.add_run(chunk), bold=False, italic=bool(i % 2))
    return p


def body(doc, text, indent=True):
    return add_runs(para_format(doc.add_paragraph(), indent=indent), text)


def heading(doc, text, level):
    style = doc.styles[f"Heading {level}"]
    p = doc.add_paragraph(style=style)
    para_format(p, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT,
                keep=True)
    set_font(p.add_run(text), bold=True, italic=False)
    return p


def centered(doc, text, bold=False):
    p = para_format(doc.add_paragraph(), indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run(text), bold=bold, italic=False)
    return p


def label(doc, kind, n, title):
    """Rótulo APA: número en negrita y título en cursiva en la línea siguiente."""
    p = para_format(doc.add_paragraph(), indent=False, keep=True)
    set_font(p.add_run(f"{kind} {n}"), bold=True, italic=False)
    p = para_format(doc.add_paragraph(), indent=False, keep=True)
    set_font(p.add_run(title), bold=False, italic=True)


def note(doc, text):
    p = para_format(doc.add_paragraph(), indent=False)
    set_font(p.add_run("Nota."), bold=False, italic=True)
    add_runs(p, " " + text)


def figure(doc, n, title, path, note_text):
    label(doc, "Figura", n, title)
    p = para_format(doc.add_paragraph(), indent=False, align=WD_ALIGN_PARAGRAPH.CENTER,
                    spacing=WD_LINE_SPACING.SINGLE, keep=True)
    p.add_run().add_picture(str(path), width=Inches(6.0))
    note(doc, note_text)


def _border(cell, edges):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if edge in edges:
            el.set(qn("w:val"), "single"), el.set(qn("w:sz"), "8"), el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tcpr.append(borders)


def table(doc, n, title, df, note_text, widths=None):
    """Tabla APA: solo líneas horizontales (arriba, bajo el encabezado y al final)."""
    label(doc, "Tabla", n, title)
    t = doc.add_table(rows=len(df) + 1, cols=len(df.columns))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(layout)
    for row in t.rows:  # filas que no se parten y tabla que no se separa de su rótulo
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
    if widths:
        for j, col in enumerate(t.columns):
            col.width = Inches(widths[j])
    for j, col in enumerate(df.columns):
        rows = [col] + [str(v) for v in df[col]]
        for i, txt in enumerate(rows):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            para_format(p, indent=False, spacing=WD_LINE_SPACING.SINGLE,
                        align=WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            set_font(p.add_run(txt), bold=False, italic=False)
            p.paragraph_format.keep_with_next = i < len(df)
            edges = {"top", "bottom"} if i == 0 else ({"bottom"} if i == len(df) else set())
            _border(cell, edges)
            if widths:
                cell.width = Inches(widths[j])
    note(doc, note_text)


def page_number_header(section):
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    set_font(run)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, Inches(1))
    for name in ("Normal", "Heading 1", "Heading 2"):
        st = doc.styles[name]
        set_font(st, bold=name != "Normal", italic=False)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        st.paragraph_format.space_before = st.paragraph_format.space_after = Pt(0)
    page_number_header(s)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- datos de outputs/
def load():
    d = {
        "resumen": pd.read_csv(OUT / "resumen_datasets.csv"),
        "variables": pd.read_csv(OUT / "variables.csv"),
        "metricas": pd.read_csv(OUT / "metricas_modelo.csv"),
        "sens": pd.read_csv(OUT / "sensibilidad_umbral.csv"),
        "res_sens": pd.read_csv(OUT / "resultados_sensibilidad.csv"),
        "final": pd.read_csv(OUT / "resumen_final.csv").set_index("modelo"),
    }
    m = d["metricas"]
    w = m.pivot_table(index=["event_id", "objetivo"], columns="modelo", values="rmse_val")
    med = m.groupby(["objetivo", "modelo"])["rmse_val"].median().unstack()
    d["gbm_wins"], d["n_combos"] = int((w["gbm"] < w["lineal"]).sum()), len(w)
    d["med_rmse"] = m.groupby("modelo")["rmse_val"].median()
    d["comp_gbm_better"] = int((med["gbm"] < med["lineal"]).sum())
    d["comp_total"] = len(med)
    d["comp_lineal_better"] = list(med.index[med["gbm"] >= med["lineal"]])
    r = d["res_sens"]
    d["main_res"] = r[r["q"] == 0.995]
    d["det"] = r[r["detectado"] & (r["tipo"] == "anomaly")]
    d["fa"] = r[r["detectado"] & (r["tipo"] == "normal")]
    return d


DESC_ES = {
    "Generator bearing failure": "Falla de rodamiento del generador",
    "Gearbox failure": "Falla de caja multiplicadora",
    "Gearbox bearings damaged": "Rodamientos de la caja dañados",
    "Hydraulic group": "Falla del grupo hidráulico",
    "Transformer failure": "Falla del transformador",
}
MODEL_ES = {"lineal": "Regresión lineal", "gbm": "Gradient boosting"}


# ---------------------------------------------------------------- contenido
def build():
    d = load()
    doc = Document()
    setup(doc)
    fin, sens, det = d["final"], d["sens"], d["det"]
    an_total, no_total = int(fin.loc["gbm", "anomalias_total"]), int(fin.loc["gbm", "normales_total"])
    rs = d["resumen"]
    counts = rs[rs["tipo"] == "anomaly"]["descripcion"].map(DESC_ES).value_counts()
    best = sens.loc[sens["CARE"].idxmax()]
    q_main = 0.995
    mr = d["main_res"]
    max_crit_main = mr[(mr["tipo"] == "anomaly") & (mr["modelo"] == "gbm")]["max_criticidad"].max()
    gb = det[det["modelo"] == "gbm"]
    # Las afirmaciones cualitativas del texto dependen de estos hechos; si los datos cambian, el build falla.
    assert set(det["event_id"]) == {0, 40, 45, 73}, set(det["event_id"])
    assert int(sens["anomalias_detectadas"].max()) * 4 == an_total
    assert d["comp_total"] - d["comp_gbm_better"] == 1
    assert (fin["anomalias_detectadas"] == 0).all() and (fin["falsas_alarmas"] == 0).all()
    alarm = det.groupby("event_id")["componente_alarma"].agg(set)
    assert alarm[0] == {"Rodamiento generador (lado acople)"} and alarm[40] == {"Aceite grupo hidráulico"}

    # ---- portada (APA 7, estudiante)
    for _ in range(3):
        centered(doc, "")
    centered(doc, TITLE, bold=True)
    centered(doc, "")
    centered(doc, AUTHOR)
    centered(doc, UNIVERSITY)
    centered(doc, COURSE)
    centered(doc, "Docentes: " + ", ".join(TEACHERS))
    today = dt.date.today()
    centered(doc, f"{today.day} de {MONTHS[today.month - 1]} de {today.year}")
    page_break(doc)

    # ---- resumen
    centered(doc, "Resumen", bold=True)
    body(doc, (
        "Este trabajo evalúa con cuánta anticipación un modelo entrenado únicamente con operación sana puede "
        "detectar la falla de un componente en un aerogenerador. Se usaron los 22 conjuntos del Parque A de CARE to "
        f"Compare ({an_total} fallas documentadas y {no_total} periodos normales). Para cada turbina se entrenaron "
        "modelos de comportamiento normal (regresión lineal y gradient boosting) que predicen siete temperaturas de "
        "componentes a partir de la potencia, la velocidad del viento, la temperatura ambiente y la velocidad del "
        "rotor; el residuo estandarizado se suavizó con EWMA y se comparó con un umbral calibrado en datos sanos. "
        f"El gradient boosting redujo el error de predicción (RMSE mediano de {fmt(d['med_rmse']['gbm'], 2)} °C "
        f"frente a {fmt(d['med_rmse']['lineal'], 2)} °C), pero con el umbral fijado a priori ninguno de los dos "
        f"modelos detectó fallas ni generó falsas alarmas (CARE = 0). Al relajar el umbral se detectaron como máximo "
        f"{int(sens['anomalias_detectadas'].max())} de {an_total} fallas, con anticipaciones entre "
        f"{fmt(det['anticipacion_dias'].min())} y {fmt(det['anticipacion_dias'].max())} días, y el mejor CARE fue "
        f"{fmt(best['CARE'], 2)}, cercano al azar. Se concluye que, en este parque, las temperaturas modeladas con "
        "esas cuatro entradas no anticipan de forma confiable la mayoría de las fallas."), indent=False)
    p = para_format(doc.add_paragraph(), indent=True)
    set_font(p.add_run("Palabras clave:"), bold=False, italic=True)
    add_runs(p, " aerogeneradores, SCADA, modelo de comportamiento normal, detección temprana de fallas, CARE")
    page_break(doc)

    # ---- introducción (APA: sin encabezado "Introducción")
    centered(doc, TITLE, bold=True)
    body(doc, (
        "La integración de energía eólica a los sistemas de potencia exige que los aerogeneradores mantengan una "
        "alta disponibilidad. Una falla no anticipada en la caja multiplicadora, el generador o el transformador "
        "detiene la turbina durante días y obliga a mantenimientos correctivos costosos. Los sistemas SCADA registran "
        "cada diez minutos la potencia, el viento y decenas de temperaturas internas, lo que permite vigilar el "
        "estado de la máquina sin instalar sensores adicionales."))
    body(doc, (
        "Un enfoque habitual es el modelo de comportamiento normal: se entrena un modelo que predice una variable "
        "del componente, como su temperatura, solo con datos de operación sana, y se interpreta una desviación "
        "persistente entre lo medido y lo predicho como señal de deterioro. Gück et al. (2024) publicaron el conjunto "
        "CARE to Compare y la métrica CARE para comparar estos métodos de forma reproducible con fallas reales "
        "documentadas."))
    body(doc, (
        "La pregunta de este trabajo es: ¿con cuánta anticipación puede un modelo entrenado únicamente con operación "
        "sana detectar una falla de componente en un aerogenerador? Para responderla se comparan una regresión "
        "lineal y un modelo de gradient boosting, se detectan anomalías en su residuo con un umbral EWMA y se evalúa "
        "el resultado contra las fallas documentadas con la anticipación, las falsas alarmas y la métrica CARE."))

    # ---- datos
    heading(doc, "Datos", 1)
    body(doc, (
        "Se usó la versión 6 de CARE to Compare (Gück y Roelofs, 2025), restringida al Parque A. Este parque tiene "
        f"22 conjuntos de datos de cinco turbinas: {an_total} terminan en una falla documentada y {no_total} son "
        "periodos normales. Cada conjunto contiene alrededor de un año de entrenamiento y un periodo de predicción "
        f"de entre {fmt(rs['filas_prediccion'].min() / 144, 0)} y {fmt(rs['filas_prediccion'].max() / 144, 0)} días, "
        "con registros cada diez minutos. Las fallas documentadas son: "
        + ", ".join(f"{k.lower()} ({v})" for k, v in counts.items()) + "."))
    body(doc, (
        "En cada conjunto de falla, los autores marcan una ventana de anomalía cuyo final corresponde al inicio de la "
        "falla. Los nombres de los sensores están anonimizados, por lo que las variables se identificaron con el "
        "archivo de descripción de sensores incluido en el conjunto (Tabla 1). La potencia viene normalizada entre 0 y 1."))
    v = d["variables"].copy()
    v["Rol"] = v["rol"].map({"entrada": "Entrada", "objetivo": "Temperatura modelada"})
    v = v.rename(columns={"nombre": "Variable", "columna": "Columna", "descripcion_original": "Descripción original"})
    table(doc, 1, "Variables del Parque A usadas en el modelo", v[["Rol", "Variable", "Columna"]],
          "Columnas del archivo de datos del Parque A identificadas con su archivo de descripción de sensores. Las "
          "entradas son las cuatro variables que pide el enunciado.", widths=[1.9, 2.7, 1.7])

    # ---- metodología
    heading(doc, "Metodología", 1)
    heading(doc, "Carga y limpieza", 2)
    body(doc, (
        "Para entrenar se conservaron solo los registros del periodo de entrenamiento con estado de operación normal, "
        "potencia positiva y sin valores faltantes. En el Parque A no existe un código de estado para potencia "
        "limitada, por lo que se excluyeron además los registros cuya potencia real fue menor que el 90 % de la "
        "potencia posible que reporta el propio SCADA. Así se filtran las paradas y la operación con potencia "
        f"limitada; se conservó entre {fmt(rs['pct_conservado'].min(), 0)} % y {fmt(rs['pct_conservado'].max(), 0)} % "
        "de los registros de entrenamiento según la turbina."))
    body(doc, (
        "Durante la inspección se encontró que, en los periodos de predicción con falla, la mayoría de los registros "
        "tiene estado de parada o mantenimiento aunque la turbina produce potencia. Los autores advierten en el "
        "repositorio que en el Parque A esas etiquetas pueden ignorarse al evaluar (Gück y Roelofs, 2025), así que en "
        "el periodo de predicción se puntuaron todos los registros con potencia positiva."))
    heading(doc, "Modelo de comportamiento normal", 2)
    body(doc, (
        "Para cada turbina se entrenó un modelo por cada una de las siete temperaturas de la Tabla 1. Se usaron las "
        "mismas siete temperaturas en todos los conjuntos para no aprovechar la etiqueta del evento y para tratar "
        "igual a los conjuntos de falla y a los normales. Las entradas fueron la potencia, la velocidad del viento, "
        "la temperatura ambiente y la velocidad del rotor, más las medias móviles de una y seis horas de la potencia y "
        "de la temperatura ambiente, que representan la inercia térmica. Se comparó una regresión lineal con un "
        "HistGradientBoostingRegressor de scikit-learn (Pedregosa et al., 2011) con sus parámetros por defecto. El "
        "último 20 % del entrenamiento, en orden temporal, se reservó como validación."))
    heading(doc, "Detección", 2)
    body(doc, (
        "El residuo de cada temperatura se estandarizó con la media y la desviación del residuo en validación y se "
        "suavizó con un EWMA de factor 0,05, equivalente a una memoria de unas tres horas. El puntaje de anomalía de "
        "cada instante es el mayor valor absoluto de los siete EWMA. Un instante se marca como anómalo si su puntaje "
        f"supera el cuantil {fmt(q_main, 3)} del puntaje en validación. Este umbral se fijó antes de ver el periodo "
        "de predicción."))
    heading(doc, "Evaluación", 2)
    body(doc, (
        "Se usó la implementación oficial de CARE publicada por los autores (Fraunhofer IEE, 2026). Un evento se "
        "considera detectado cuando un contador de criticidad, que sube con cada instante anómalo y baja con cada "
        "instante normal, alcanza 72, es decir, unas 12 horas netas de anomalía. La anticipación es el tiempo entre "
        "esa primera detección y el inicio de la falla. Una detección en un conjunto normal es una falsa alarma. CARE "
        "combina la cobertura, la exactitud en conjuntos normales, la confiabilidad por evento y la anticipación; "
        "vale 0 si no se detecta ninguna falla. Como complemento se repitió la evaluación con cuantiles de umbral "
        "entre 0,90 y 0,999. Este análisis de sensibilidad se reporta por separado, porque elegir el umbral mirando "
        "el periodo de predicción sobrestimaría el desempeño."))

    # ---- resultados
    heading(doc, "Resultados", 1)
    heading(doc, "Desempeño de los modelos de comportamiento normal", 2)
    lin_better = ", ".join(c for c in d["comp_lineal_better"])
    names = dict(zip(d["variables"]["columna"], d["variables"]["nombre"]))
    body(doc, (
        f"El gradient boosting obtuvo menor RMSE de validación que la regresión lineal en {d['gbm_wins']} de las "
        f"{d['n_combos']} combinaciones de turbina y componente. La Figura 1 muestra que su RMSE mediano es menor en "
        f"{d['comp_gbm_better']} de los {d['comp_total']} componentes; en el restante "
        f"({names.get(lin_better, lin_better).lower().replace(' (', ', ').replace(')', '')}) ambos modelos quedan prácticamente iguales. La mayor mejora "
        "aparece en las temperaturas de la caja multiplicadora, del devanado y del grupo hidráulico."))
    figure(doc, 1, "Error de predicción en validación por componente y modelo", OUT / "fig1_rmse_modelos.png",
           f"Cada punto es la mediana del RMSE sobre los 22 conjuntos del Parque A. RMSE = raíz del error "
           "cuadrático medio.")
    body(doc, (
        "La Figura 2 ilustra la diferencia en una turbina con falla: en los últimos días de validación, la regresión "
        "lineal subestima los picos diarios de temperatura del aceite de la caja, mientras que el gradient boosting "
        "los sigue de cerca."))
    figure(doc, 2, "Temperatura real y predicha del aceite de la caja multiplicadora en validación (evento 0)",
           OUT / "fig2_real_vs_predicho.png",
           "Últimos cinco días de validación del conjunto 0. Los cortes de las líneas corresponden a registros "
           "excluidos por parada o potencia limitada.")

    heading(doc, "Detección de fallas con el umbral principal", 2)
    body(doc, (
        f"Con el umbral del cuantil {fmt(q_main, 3)}, ninguno de los dos modelos detectó alguna de las {an_total} "
        f"fallas y ninguno generó falsas alarmas en los {no_total} conjuntos normales; por lo tanto, CARE = 0 en "
        "ambos casos. La Figura 3 muestra por qué: en las doce ventanas de falla el puntaje del gradient boosting cruza "
        "el umbral solo en picos breves y vuelve a bajar, de modo que el contador de criticidad nunca se acerca a 72 "
        f"(el máximo alcanzado fue {int(max_crit_main)}). El evento 40, una falla de rodamiento del generador, presenta "
        "los cruces más altos, sobre todo en los cinco días previos a la falla, pero también son intermitentes."))
    figure(doc, 3, "Puntaje de anomalía del gradient boosting en el periodo de predicción de cada falla",
           OUT / "fig3_puntaje_fallas.png",
           "El puntaje se divide por el umbral de cada turbina, de modo que el umbral vale 1 en todos los paneles. "
           "Los valores mayores que 3 se recortan. Los huecos corresponden a registros sin potencia.")

    heading(doc, "Sensibilidad al umbral", 2)
    body(doc, (
        "La Figura 4 y la Tabla 2 muestran qué ocurre al relajar el umbral. Bajar el cuantil aumenta las detecciones, "
        f"pero solo hasta {int(sens['anomalias_detectadas'].max())} de {an_total} fallas, y a la vez aparecen falsas "
        f"alarmas (hasta {int(sens['falsas_alarmas'].max())} de {no_total}). CARE nunca supera "
        f"{fmt(sens['CARE'].max(), 2)}: queda en torno al valor de un detector aleatorio (0,5) y por debajo del "
        "0,66 que obtuvo el autoencoder de referencia de Gück et al. (2024)."))
    figure(doc, 4, "Efecto del umbral sobre las detecciones, las falsas alarmas y CARE",
           OUT / "fig4_sensibilidad_umbral.png",
           "La línea vertical marca el umbral principal, fijado a priori. Los valores de referencia 0,5 (azar) y "
           "0,66 (autoencoder) provienen de Gück et al. (2024) y se calcularon sobre los 95 conjuntos de los tres "
           "parques, no solo sobre el Parque A.")
    t = sens.copy()
    t["Modelo"] = t["modelo"].map(MODEL_ES)
    t["Cuantil"] = t["q"].map(lambda x: f"{x:g}".replace(".", ","))
    t["Fallas detectadas"] = t["anomalias_detectadas"].astype(int).astype(str) + f" de {an_total}"
    t["Falsas alarmas"] = t["falsas_alarmas"].astype(int).astype(str) + f" de {no_total}"
    t["Anticipación mediana (días)"] = t["anticipacion_mediana_dias"].map(lambda x: "n.a." if pd.isna(x) else fmt(x))
    t["CARE"] = t["CARE"].map(lambda x: fmt(x, 2))
    table(doc, 2, "Resultados por modelo y cuantil del umbral",
          t[["Modelo", "Cuantil", "Fallas detectadas", "Falsas alarmas", "Anticipación mediana (días)", "CARE"]],
          "Fila con cuantil 0,995: umbral principal, fijado a priori. Las demás filas son análisis de sensibilidad. "
          "n.a. = no aplica porque no hubo detecciones.", widths=[1.6, 0.8, 1.1, 1.0, 1.2, 0.7])

    dd = det.copy()
    dd = dd.sort_values(["event_id", "modelo", "q"]).groupby(["event_id", "modelo"]).last().reset_index()
    dd["Evento"] = dd["event_id"].astype(str)
    dd["Falla documentada"] = dd["descripcion"].map(DESC_ES)
    dd["Modelo"] = dd["modelo"].map(MODEL_ES)
    dd["Cuantil"] = dd["q"].map(lambda x: f"{x:g}".replace(".", ","))
    dd["Anticipación (días)"] = dd["anticipacion_dias"].map(fmt)
    dd["Temperatura que disparó la alarma"] = dd["componente_alarma"]
    body(doc, (
        "La Tabla 3 detalla las fallas detectadas en el análisis de sensibilidad. Solo en el evento 0 la alarma la "
        "dispara la temperatura del componente que falla, el rodamiento del generador, y lo hace con una anticipación "
        f"de {fmt(dd[dd['event_id'] == 0]['anticipacion_dias'].max())} días. En el evento 40 la detección llega "
        f"{fmt(dd[dd['event_id'] == 40]['anticipacion_dias'].max())} días antes de la falla, pero la dispara el aceite "
        "del grupo hidráulico, que no es el componente dañado; en los eventos 45 y 73 ocurre lo mismo con componentes "
        "que tampoco corresponden a la falla."))
    table(doc, 3, "Fallas detectadas en el análisis de sensibilidad",
          dd[["Evento", "Falla documentada", "Modelo", "Cuantil", "Anticipación (días)",
              "Temperatura que disparó la alarma"]],
          "Para cada evento y modelo se muestra el cuantil más alto con el que se detectó la falla. La anticipación "
          "se mide desde la primera detección hasta el inicio de la falla.", widths=[0.75, 1.65, 1.15, 0.7, 1.0, 1.35])

    # ---- discusión
    heading(doc, "Discusión", 1)
    heading(doc, "Respuesta a la pregunta", 2)
    body(doc, (
        "Con el método planteado y un umbral fijado sin mirar las fallas, un modelo entrenado solo con operación "
        f"sana no anticipó ninguna de las {an_total} fallas del Parque A, aunque tampoco produjo falsas alarmas. Al "
        "relajar el umbral, la anticipación observada fue de entre "
        f"{fmt(det['anticipacion_dias'].min())} y {fmt(det['anticipacion_dias'].max())} días. En la mejor "
        f"configuración ({MODEL_ES[best['modelo']].lower()}, cuantil {fmt(best['q'], 2)}), un modelo entrenado solo "
        f"con operación sana detecta fallas con una mediana de {fmt(best['anticipacion_mediana_dias'])} días de "
        f"anticipación ({int(best['anomalias_detectadas'])}/{an_total} fallas detectadas, "
        f"{int(best['falsas_alarmas'])} {'falsa alarma' if int(best['falsas_alarmas']) == 1 else 'falsas alarmas'}, CARE = {fmt(best['CARE'], 2)}). Ese resultado se eligió "
        "después de ver el periodo de predicción, así que es una cota optimista. La única detección físicamente "
        f"coherente con el componente dañado ocurrió entre {fmt(gb[gb['event_id'] == 0]['anticipacion_dias'].min())} y "
        f"{fmt(gb[gb['event_id'] == 0]['anticipacion_dias'].max())} "
        "días antes de la falla."))
    body(doc, (
        "El mejor ajuste del gradient boosting no se tradujo en mejor detección. Un modelo más preciso reduce el "
        "ruido del residuo, pero si la falla no altera la temperatura modelada, el residuo no cambia. La mitad de "
        "las fallas del Parque A son del grupo hidráulico, cuya degradación probablemente se manifiesta en presión o "
        "en eventos de control más que en la temperatura del aceite; la Figura 3 muestra que en esos seis eventos el "
        "puntaje permanece bajo durante casi toda la ventana."))
    heading(doc, "Limitaciones", 2)
    body(doc, (
        "El análisis cubre solo el Parque A, con doce fallas, por lo que las proporciones tienen gran incertidumbre. "
        "Las entradas se limitaron a las cuatro variables del enunciado; otras señales, como la temperatura de la "
        "góndola o las variables eléctricas, podrían explicar mejor la temperatura de los componentes. No se ajustaron "
        "hiperparámetros, el umbral de criticidad de 72 registros es exigente para ventanas de una semana y las "
        "etiquetas de estado del periodo de predicción no son confiables en este parque. Por último, la comparación "
        "con el autoencoder de referencia es solo orientativa, porque su CARE se calculó sobre los tres parques."))

    # ---- conclusiones
    heading(doc, "Conclusiones", 1)
    body(doc, (
        "Se implementó de forma reproducible un modelo de comportamiento normal para siete temperaturas de "
        "componentes del Parque A de CARE to Compare, con detección EWMA y evaluación mediante la implementación "
        "oficial de CARE. El gradient boosting predice la temperatura mejor que la regresión lineal, pero ninguno de "
        "los dos anticipa de forma confiable las fallas documentadas: con el umbral fijado a priori no hay detecciones, "
        "y al relajarlo se detecta como máximo una de cada cuatro fallas, con un CARE cercano al azar. La "
        "anticipación de varios días que reporta la literatura no se obtiene en este parque con temperaturas "
        "explicadas solo por la potencia, el viento, la temperatura ambiente y la velocidad del rotor. Un trabajo "
        "futuro debería incorporar más variables del SCADA, como hace el autoencoder de referencia, y evaluar los tres "
        "parques."))

    # ---- referencias
    page_break(doc)
    centered(doc, "Referencias", bold=True)
    refs = [
        "Fraunhofer Institute for Energy Economics and Energy System Technology. (2026). *EnergyFaultDetector* "
        "(Versión 0.8.1) [Software]. https://github.com/AEFDI/EnergyFaultDetector",
        "Gück, C., y Roelofs, C. (2025). *Wind turbine SCADA data for early fault detection* (Versión v6) "
        "[Conjunto de datos]. Zenodo. https://doi.org/10.5281/zenodo.15846963",
        "Gück, C., Roelofs, C. M. A., y Faulstich, S. (2024). CARE to Compare: A real-world benchmark dataset for "
        "early fault detection in wind turbine data. *Data, 9*(12), 138. https://doi.org/10.3390/data9120138",
        "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., Blondel, M., "
        "Prettenhofer, P., Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., Cournapeau, D., Brucher, M., "
        "Perrot, M., y Duchesnay, É. (2011). Scikit-learn: Machine learning in Python. *Journal of Machine Learning "
        "Research, 12*, 2825-2830.",
    ]
    for ref in refs:
        add_runs(para_format(doc.add_paragraph(), indent=False, hanging=True), ref)

    doc.save(DOCX)
    return DOCX


# ---------------------------------------------------------------- verificación automática del formato
def verify(path=DOCX):
    doc = Document(path)
    problems = []
    s = doc.sections[0]
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        if abs(getattr(s, side) - Inches(1)) > Inches(0.01):
            problems.append(f"margen {side} distinto de 2,54 cm")
    header_xml = s.header._element.xml
    if "PAGE" not in header_xml:
        problems.append("falta número de página en el encabezado")

    paragraphs = list(doc.paragraphs) + [p for t in doc.tables for row in t.rows for c in row.cells
                                         for p in c.paragraphs] + list(s.header.paragraphs)
    table_pars = {id(p._p) for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs}
    header_pars = {id(p._p) for p in s.header.paragraphs}
    for p in paragraphs:
        for ch in FORBIDDEN:
            if ch in p.text:
                problems.append(f"guion prohibido en: {p.text[:60]}")
        for r in p.runs:
            if not r.text:
                continue
            fonts = r._r.rPr.find(qn("w:rFonts")) if r._r.rPr is not None else None
            if r.font.name != FONT or fonts is None or fonts.get(qn("w:eastAsia")) != FONT:
                problems.append(f"fuente distinta de Arial: {r.text[:40]}")
            if r.font.size != SIZE:
                problems.append(f"tamaño distinto de 11: {r.text[:40]}")
            if r.font.color.rgb not in (None, BLACK):
                problems.append(f"texto con color: {r.text[:40]}")
        has_text = any(r.text for r in p.runs)
        if has_text and id(p._p) not in table_pars and id(p._p) not in header_pars:
            if p.paragraph_format.line_spacing_rule != WD_LINE_SPACING.DOUBLE:
                problems.append(f"interlineado no doble: {p.text[:40]}")
    for st in ("Normal", "Heading 1", "Heading 2"):
        if doc.styles[st].font.color.rgb not in (None, BLACK):
            problems.append(f"estilo {st} con color")
    text = "\n".join(p.text for p in doc.paragraphs)
    for item in (AUTHOR, UNIVERSITY, COURSE, *TEACHERS):
        if item not in text:
            problems.append(f"portada sin: {item}")
    bold_runs = {r.text for p in doc.paragraphs for r in p.runs if r.bold}
    return problems, sorted(bold_runs), len(doc.inline_shapes)


if __name__ == "__main__":
    out = build()
    problems, bold, n_figs = verify(out)
    print(f"Generado: {out}")
    print(f"Figuras insertadas: {n_figs}")
    print("Textos en negrita:", bold)
    print("Problemas de formato:", problems or "ninguno")
